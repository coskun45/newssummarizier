"""
Tests for app/api/routes/bulletin.py.
"""
import asyncio
import io
import re
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.core.config import settings as app_settings
from app.db import crud, models
from app.services import bulletin_service, summary_service
from tests.conftest import _make_article, _make_feed, _make_summary

# The autouse stub below replaces it on the module; its own test needs the real one.
_real_group_category_articles = bulletin_service.group_category_articles
_real_classify_articles_for_bulletin = bulletin_service.classify_articles_for_bulletin
_real_pick_bulletin_highlights = bulletin_service.pick_bulletin_highlights


@pytest.fixture(autouse=True)
def _isolated_bulletin_storage(tmp_path, monkeypatch):
    """Every /generate call now writes a .docx to BULLETIN_STORAGE_DIR — redirect
    it to a pytest tmp dir so tests never write into the real repo checkout."""
    monkeypatch.setattr(app_settings, "bulletin_storage_dir", str(tmp_path))


def _make_bulletin_category(db_session, name="AVRUPA", display_order=0):
    category = models.BulletinCategory(name=name, display_order=display_order)
    db_session.add(category)
    db_session.commit()
    db_session.refresh(category)
    return category


# ==================== GET /categories (list) ====================

def test_list_categories_requires_auth(client):
    response = client.get("/api/bulletin/categories")
    assert response.status_code == 401


def test_list_categories_empty(client, auth_headers):
    response = client.get("/api/bulletin/categories", headers=auth_headers)
    assert response.status_code == 200
    assert response.json() == []


def test_list_categories_returns_display_order(client, auth_headers, db_session):
    _make_bulletin_category(db_session, name="AMERİKA", display_order=1)
    _make_bulletin_category(db_session, name="AVRUPA", display_order=0)

    response = client.get("/api/bulletin/categories", headers=auth_headers)

    assert response.status_code == 200
    assert [c["name"] for c in response.json()] == ["AVRUPA", "AMERİKA"]


# ==================== POST /categories (create) ====================

def test_create_category_requires_admin(client, auth_headers):
    response = client.post("/api/bulletin/categories", json={"name": "AVRUPA"}, headers=auth_headers)
    assert response.status_code == 403


def test_create_category_success(client, admin_headers):
    response = client.post("/api/bulletin/categories", json={"name": "AVRUPA"}, headers=admin_headers)
    assert response.status_code == 200
    assert response.json()["name"] == "AVRUPA"


def test_create_category_rejects_duplicate_name(client, admin_headers, db_session):
    _make_bulletin_category(db_session, name="AVRUPA")

    response = client.post("/api/bulletin/categories", json={"name": "AVRUPA"}, headers=admin_headers)

    assert response.status_code == 400


def test_create_category_rejects_empty_name(client, admin_headers):
    response = client.post("/api/bulletin/categories", json={"name": "   "}, headers=admin_headers)
    assert response.status_code == 400


def test_create_category_rejects_case_insensitive_duplicate(client, admin_headers, db_session):
    """"Avrupa" and "AVRUPA" fold to the same key in
    docx_service._normalize() and would collide onto the same template
    heading when rendering a bulletin — the exact-match check let both be
    created. Regression test for get_bulletin_category_by_name's
    case/Turkish-I-insensitive lookup."""
    _make_bulletin_category(db_session, name="Avrupa")

    response = client.post("/api/bulletin/categories", json={"name": "AVRUPA"}, headers=admin_headers)

    assert response.status_code == 400


# ==================== PUT /categories/{id} (update) ====================

def test_update_category_requires_admin(client, auth_headers, db_session):
    category = _make_bulletin_category(db_session)
    response = client.put(
        f"/api/bulletin/categories/{category.id}", json={"name": "AMERİKA"}, headers=auth_headers
    )
    assert response.status_code == 403


def test_update_category_renames(client, admin_headers, db_session):
    category = _make_bulletin_category(db_session, name="AVRUPA")

    response = client.put(
        f"/api/bulletin/categories/{category.id}", json={"name": "AMERİKA"}, headers=admin_headers
    )

    assert response.status_code == 200
    assert response.json()["name"] == "AMERİKA"


def test_update_category_404_when_missing(client, admin_headers):
    response = client.put("/api/bulletin/categories/999999", json={"name": "X"}, headers=admin_headers)
    assert response.status_code == 404


def test_update_category_rejects_case_insensitive_duplicate(client, admin_headers, db_session):
    _make_bulletin_category(db_session, name="Avrupa", display_order=0)
    other = _make_bulletin_category(db_session, name="AMERİKA", display_order=1)

    response = client.put(
        f"/api/bulletin/categories/{other.id}", json={"name": "avrupa"}, headers=admin_headers
    )

    assert response.status_code == 400


# ==================== DELETE /categories/{id} ====================

def test_delete_category_requires_admin(client, auth_headers, db_session):
    category = _make_bulletin_category(db_session)
    response = client.delete(f"/api/bulletin/categories/{category.id}", headers=auth_headers)
    assert response.status_code == 403


def test_delete_category_success(client, admin_headers, db_session):
    category = _make_bulletin_category(db_session)
    response = client.delete(f"/api/bulletin/categories/{category.id}", headers=admin_headers)
    assert response.status_code == 200


def test_delete_category_404_when_missing(client, admin_headers):
    response = client.delete("/api/bulletin/categories/999999", headers=admin_headers)
    assert response.status_code == 404


# ==================== PUT /categories/reorder ====================

def test_reorder_categories_requires_admin(client, auth_headers, db_session):
    category = _make_bulletin_category(db_session)
    response = client.put(
        "/api/bulletin/categories/reorder", json={"ordered_ids": [category.id]}, headers=auth_headers
    )
    assert response.status_code == 403


def test_reorder_categories_success(client, admin_headers, db_session):
    a = _make_bulletin_category(db_session, name="AVRUPA", display_order=0)
    b = _make_bulletin_category(db_session, name="AMERİKA", display_order=1)

    response = client.put(
        "/api/bulletin/categories/reorder",
        json={"ordered_ids": [b.id, a.id]},
        headers=admin_headers,
    )

    assert response.status_code == 200
    assert [c["name"] for c in response.json()] == ["AMERİKA", "AVRUPA"]


# ==================== POST /generate ====================

def test_generate_bulletin_requires_auth(client):
    response = client.post("/api/bulletin/generate", json={})
    assert response.status_code == 401


def test_generate_bulletin_no_categories_returns_400(client, auth_headers):
    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)
    assert response.status_code == 400


def test_generate_bulletin_no_articles_returns_400(client, auth_headers, db_session):
    _make_bulletin_category(db_session, name="AVRUPA")

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 400


def test_generate_bulletin_rejects_invalid_priority(client, auth_headers, db_session):
    _make_bulletin_category(db_session, name="AVRUPA")

    response = client.post(
        "/api/bulletin/generate", json={"priorities": ["urgent"]}, headers=auth_headers
    )

    assert response.status_code == 400


async def _fake_classify(db, articles, category_names, blocks=None):
    if not articles:
        return []
    return [
        {"article_id": a.id, "top_category": category_names[0], "subcategory": "Test Alt Başlık", "type": "haber"}
        for a in articles
    ]


async def _no_highlights(articles, blocks, limit=bulletin_service.HIGHLIGHT_LIMIT):
    return []


async def _single_group(category_name, articles, blocks):
    return [{"title": "TEST ALT BAŞLIK", "article_ids": [a.id for a in articles]}] if articles else []


@pytest.fixture(autouse=True)
def _stub_bulletin_llm(monkeypatch):
    """No test calls OpenAI: classification, highlights, grouping and the brief summarizer are
    stubbed; tests override the one they exercise. Returns the brief-summary call log."""
    calls = []

    async def fake_generate_summary(title, content, summary_type="standard", source=None, author_hint=None):
        calls.append({"title": title, "summary_type": summary_type, "source": source})
        return {
            "summary_text": f"📌 **{source} - {title}**\n🔹 {title} özeti.",
            "model_used": "gpt-test", "tokens_used": 10, "cost": 0.001,
            "author": None, "structured": True,
        }

    monkeypatch.setattr(bulletin_service, "classify_articles_for_bulletin", _fake_classify)
    monkeypatch.setattr(bulletin_service, "pick_bulletin_highlights", _no_highlights)
    monkeypatch.setattr(bulletin_service, "group_category_articles", _single_group)
    async def no_cost_limits():
        return None

    monkeypatch.setattr(summary_service, "generate_summary", fake_generate_summary)
    # check_cost_limits opens its own SessionLocal (the real DB), not the test session.
    monkeypatch.setattr(bulletin_service, "check_cost_limits", no_cost_limits)
    return calls


def _generate(client, headers, **body):
    response = client.post("/api/bulletin/generate", json=body, headers=headers)
    assert response.status_code == 200, response.text
    return response


def _paragraphs(response):
    return DocxDocument(io.BytesIO(response.content)).paragraphs


def test_generate_bulletin_success(client, auth_headers, db_session):
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(
        db_session, feed.id, title="Test Haberi", priority="high",
        published_at=now - timedelta(hours=1),
    )

    response = _generate(client, auth_headers)

    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.content[:4] == b"PK\x03\x04"

    paragraphs = _paragraphs(response)
    by_text = {p.text.strip(): p for p in paragraphs if p.text.strip()}
    # Like the template edition: category -> Heading1, topic group -> Heading3 (the TOC field
    # maps "Heading 3" to level 3); no Heading2 anywhere.
    assert by_text["AVRUPA"].style.style_id == "Heading1"
    assert by_text["TEST ALT BAŞLIK"].style.style_id == "Heading3"
    assert not [p for p in paragraphs if p.style.style_id == "Heading2"]

    # The article block comes from the brief summary: bold "📌 Kaynak - Başlık" header (the
    # summarizer's "**" markers stripped) and plain "🔹" bullets.
    header = by_text["📌 Test Feed - Test Haberi"]
    assert header.runs[0].bold
    bullet = by_text["🔹 Test Haberi özeti."]
    assert not bullet.runs[0].bold
    assert not any("**" in p.text for p in paragraphs)


def test_generate_bulletin_keeps_template_sources_and_adds_rss_feeds(client, auth_headers, db_session):
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session, title="DW Türkçe")
    _make_feed(db_session, url="https://example.com/passive", title="Pasif Feed", is_active=False)
    _make_article(db_session, feed.id, title="Haber", published_at=datetime.now(timezone.utc) - timedelta(hours=1))

    texts = [p.text for p in _paragraphs(_generate(client, auth_headers))]

    assert any(t.startswith("ABD: Associated Press") for t in texts)
    rss_line = next(t for t in texts if t.startswith("RSS Beslemeleri:"))
    assert "DW Türkçe" in rss_line
    assert "Pasif Feed" not in rss_line


def test_generate_bulletin_toc_is_prerendered_with_bookmark_links(client, auth_headers, db_session):
    """Word used to be told to refresh fields on open (updateFields), which asks the reader for
    permission every time; the İçindekiler is now written out with one link per heading."""
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    _make_article(db_session, feed.id, title="Haber", published_at=datetime.now(timezone.utc) - timedelta(hours=1))

    response = _generate(client, auth_headers)

    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
        settings_xml = zf.read("word/settings.xml").decode("utf-8")
    assert "updateFields" not in settings_xml

    toc = document_xml[document_xml.index("<w:sdt>"):document_xml.index("</w:sdt>")]
    anchors = re.findall(r'w:anchor="([^"]+)"', toc)
    entry_texts = re.findall(r"<w:t[^>]*>([^<]*)</w:t>", toc)
    assert entry_texts == ["AVRUPA", "TEST ALT BAŞLIK"]
    body = document_xml[document_xml.index("</w:sdt>"):]
    for anchor in anchors:
        assert f'w:name="{anchor}"' in body
    assert toc.count('w:fldCharType="begin"') == 1
    assert toc.count('w:fldCharType="end"') == 1


def test_generate_bulletin_highlights_are_not_repeated_in_categories(
    client, auth_headers, db_session, monkeypatch
):
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    articles = [
        _make_article(db_session, feed.id, title=f"Haber {i}", published_at=now - timedelta(hours=1))
        for i in range(7)
    ]

    async def highlights(candidates, blocks, limit=bulletin_service.HIGHLIGHT_LIMIT):
        return [{"title": "BM GENEL KURULU", "article_ids": [articles[0].id, articles[1].id]}]

    monkeypatch.setattr(bulletin_service, "pick_bulletin_highlights", highlights)

    paragraphs = _paragraphs(_generate(client, auth_headers))
    texts = [p.text for p in paragraphs]
    heading1 = [p.text for p in paragraphs if p.style.style_id == "Heading1"]

    assert heading1[0] == "ÖNE ÇIKAN BAŞLIKLAR"
    assert texts[texts.index("ÖNE ÇIKAN BAŞLIKLAR") + 1] == "BM GENEL KURULU"
    assert sum("Haber 0" in t for t in texts if t.startswith("📌")) == 1
    avrupa_idx = texts.index("AVRUPA")
    assert not any("Haber 0" in t or "Haber 1" in t for t in texts[avrupa_idx:])
    assert any("Haber 6" in t for t in texts[avrupa_idx:])


def test_clean_groups_caps_highlights_at_limit_and_drops_unknown_ids():
    groups = bulletin_service._clean_groups(
        [
            {"title": "bm genel kurulu", "article_ids": [1, 2, 99, 2]},
            {"title": "Fon krizi", "article_ids": ["3", 4, 5, 6]},
        ],
        allowed_ids=[1, 2, 3, 4, 5, 6],
        limit=5,
    )
    assert groups == [
        {"title": "BM GENEL KURULU", "article_ids": [1, 2]},
        {"title": "FON KRİZİ", "article_ids": [3, 4, 5]},
    ]


def test_group_category_articles_puts_leftovers_into_other_group(monkeypatch):
    """Articles the model forgot must still appear, under the catch-all group, which goes last."""
    async def fake_completion(system_prompt, user_prompt, max_completion_tokens):
        return {"groups": [{"title": "UKRAYNA SAVAŞI", "article_ids": [1]}]}

    async def no_limits():
        return None

    monkeypatch.setattr(bulletin_service, "_call_json_completion", fake_completion)
    monkeypatch.setattr(bulletin_service, "check_cost_limits", no_limits)
    articles = [models.Article(id=i, title=f"Haber {i}", url=f"https://x/{i}") for i in (1, 2, 3)]
    blocks = {a.id: {"header": f"📌 X - {a.title}", "bullets": []} for a in articles}

    groups = asyncio.run(_real_group_category_articles("Avrupa", articles, blocks))

    assert groups == [
        {"title": "UKRAYNA SAVAŞI", "article_ids": [1]},
        {"title": "AVRUPA: DİĞER GELİŞMELER", "article_ids": [2, 3]},
    ]


def test_generate_bulletin_creates_and_saves_missing_brief_summaries(
    client, auth_headers, db_session, _stub_bulletin_llm
):
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session, title="BBC Türkçe")
    now = datetime.now(timezone.utc)
    with_brief = _make_article(db_session, feed.id, title="Özetli", published_at=now - timedelta(hours=1))
    _make_summary(db_session, with_brief.id, summary_type="brief",
                  summary_text="📌 BBC Türkçe - Özetli\n🔹 Kayıtlı kısa özet.")
    without_brief = _make_article(db_session, feed.id, title="Özetsiz", cleaned_content="Metin",
                                  published_at=now - timedelta(hours=1))

    texts = [p.text for p in _paragraphs(_generate(client, auth_headers))]

    assert [c["title"] for c in _stub_bulletin_llm] == ["Özetsiz"]
    assert _stub_bulletin_llm[0]["summary_type"] == "brief"
    assert _stub_bulletin_llm[0]["source"] == "BBC Türkçe"
    stored = db_session.query(models.Summary).filter_by(article_id=without_brief.id, summary_type="brief").all()
    assert len(stored) == 1
    assert "🔹 Kayıtlı kısa özet." in texts
    assert "🔹 Özetsiz özeti." in texts

    # Saved: the next report doesn't summarize again.
    _generate(client, auth_headers)
    assert len(_stub_bulletin_llm) == 1


def test_generate_bulletin_marks_opinion_pieces_with_yorum_icon(client, auth_headers, db_session, monkeypatch):
    async def classify_as_yorum(db, articles, category_names, blocks=None):
        return [{"article_id": a.id, "top_category": category_names[0], "type": "yorum"} for a in articles]

    monkeypatch.setattr(bulletin_service, "classify_articles_for_bulletin", classify_as_yorum)
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    _make_article(db_session, feed.id, title="Köşe Yazısı", published_at=datetime.now(timezone.utc) - timedelta(hours=1))

    texts = [p.text for p in _paragraphs(_generate(client, auth_headers))]

    assert "⭕️ Test Feed - Köşe Yazısı" in texts


def test_article_block_parses_single_line_brief():
    article = models.Article(id=1, title="Başlık", url="https://x/1")
    block = bulletin_service._article_block(article, "**📌 DW / Ali Veli - Başlık** 🔹 Bir. 🔹 İki.", "haber")
    assert block["header"] == "📌 DW / Ali Veli - Başlık"
    assert block["bullets"] == ["🔹 Bir.", "🔹 İki."]


def test_generate_bulletin_strips_html_from_raw_content_fallback(client, auth_headers, db_session, monkeypatch):
    """Regression test: when an article has no "brief" summary (and generating one fails), the
    bulletin falls back to raw_content, which for several feeds (e.g. Aydinlik) is an HTML
    fragment straight from the RSS <description> rather than plain text. A real generated
    bulletin showed this markup — <a>/<img>/<h4> tags and &#039;-style entities — rendered
    verbatim as visible text in the Word document."""
    async def failing_summary(**kwargs):
        raise RuntimeError("summarizer down")

    monkeypatch.setattr(summary_service, "generate_summary", failing_summary)
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(
        db_session, feed.id, title="HTML Kirli Haber",
        raw_content=(
            '<a href="https://example.com/haber">'
            '<img align="right" border="0" height="84" src="https://img.example.com/x.jpg" width="150" />'
            '</a>\n<h4>Firari şüpheli &#039;yakalandı&#039;</h4>'
        ),
        published_at=now - timedelta(hours=1),
    )

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))
    paragraph_texts = [p.text for p in document.paragraphs if p.text.strip()]

    assert not any("<a href" in t or "<img " in t or "<h4>" in t for t in paragraph_texts)
    assert any("Firari şüpheli 'yakalandı'" in t for t in paragraph_texts)


def test_generate_bulletin_falls_back_to_stored_summary_when_brief_fails(
    client, auth_headers, db_session, monkeypatch
):
    """A failed brief uses the stored standard summary (else the detailed one) instead of the
    plain-content snippet, and that fallback is not saved as a brief."""
    async def failing_summary(**kwargs):
        raise RuntimeError("summarizer down")

    monkeypatch.setattr(summary_service, "generate_summary", failing_summary)
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    both = _make_article(db_session, feed.id, title="Standartlı", cleaned_content="Ham metin A",
                         published_at=now - timedelta(hours=1))
    _make_summary(db_session, both.id, summary_type="detailed",
                  summary_text="📌 Test Feed - Standartlı\n\nDetaylı paragraf A.")
    _make_summary(db_session, both.id, summary_type="standard",
                  summary_text="📌 Test Feed - Standartlı\n\nStandart paragraf A.")
    detailed_only = _make_article(db_session, feed.id, title="Detaylı", cleaned_content="Ham metin B",
                                  published_at=now - timedelta(hours=1))
    _make_summary(db_session, detailed_only.id, summary_type="detailed",
                  summary_text="📌 Test Feed - Detaylı\n\nDetaylı paragraf B.")
    no_summary = _make_article(db_session, feed.id, title="Özetsiz", cleaned_content="Ham metin C",
                               published_at=now - timedelta(hours=1))

    texts = [p.text for p in _paragraphs(_generate(client, auth_headers))]

    assert "🔹 Standart paragraf A." in texts
    assert not any("Detaylı paragraf A." in t or "Ham metin A" in t for t in texts)
    assert "🔹 Detaylı paragraf B." in texts
    assert not any("Ham metin B" in t for t in texts)
    assert "🔹 Ham metin C" in texts
    stored_briefs = db_session.query(models.Summary).filter(
        models.Summary.article_id.in_([both.id, detailed_only.id, no_summary.id]),
        models.Summary.summary_type == "brief",
    ).count()
    assert stored_briefs == 0


def test_pick_bulletin_highlights_sends_only_top_ranked_candidates(monkeypatch):
    """Only the 10 highest-priority, newest articles are offered to the highlights model."""
    prompts = []

    async def fake_completion(system_prompt, user_prompt, max_completion_tokens):
        prompts.append(user_prompt)
        return {"highlights": []}

    async def no_limits():
        return None

    monkeypatch.setattr(bulletin_service, "_call_json_completion", fake_completion)
    monkeypatch.setattr(bulletin_service, "check_cost_limits", no_limits)
    now = datetime.now(timezone.utc)
    # ids 1-5 high, 6-15 med (6 newest), 16-20 low: the cap keeps 1-5 and 6-10.
    articles = [
        models.Article(
            id=i, title=f"Haber {i}", url=f"https://x/{i}",
            priority="high" if i <= 5 else "med" if i <= 15 else "low",
            published_at=now - timedelta(minutes=i),
        )
        for i in range(20, 0, -1)
    ]
    blocks = {a.id: {"header": f"📌 X - {a.title}", "bullets": []} for a in articles}

    asyncio.run(_real_pick_bulletin_highlights(articles, blocks))

    sent_ids = [int(line.split(" | ")[0]) for line in prompts[0].splitlines() if " | " in line]
    assert sent_ids == list(range(1, 11))


def test_strip_html_preserves_literal_angle_bracket_text():
    """Regression test: a real HTML parser (not a "<[^>]+>" regex) must not mistake
    literal bracketed text that isn't markup for a tag and delete it."""
    result = bulletin_service._strip_html("büyüme <10 yaş grubunda> daha hızlı oldu")
    assert "10 yaş grubunda" in result


def test_strip_html_collapses_whitespace_from_entities():
    """Regression test: entities that decode to whitespace (e.g. repeated &nbsp;)
    must be unescaped before whitespace collapsing runs, or they leak through
    uncollapsed since at collapse-time they're still literal "&nbsp;" text."""
    result = bulletin_service._strip_html("Ankara&nbsp;&nbsp;&nbsp;Istanbul")
    assert result == "Ankara Istanbul"


def test_generate_bulletin_uses_only_user_defined_categories(client, auth_headers, db_session, monkeypatch):
    """A category set unrelated to the bundled template's default region
    headings must fully replace them in the output — regression test for a
    reported issue where the generated report appeared to ignore
    user-created categories."""
    async def classify_into_custom_categories(db, articles, category_names, blocks=None):
        return [
            {"article_id": a.id, "top_category": category_names[0], "subcategory": "Alt Başlık", "type": "haber"}
            for a in articles
        ]

    monkeypatch.setattr(bulletin_service, "classify_articles_for_bulletin", classify_into_custom_categories)

    _make_bulletin_category(db_session, name="Teknoloji")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Yapay Zeka Haberi", published_at=now - timedelta(hours=1))

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))
    paragraph_texts = [p.text for p in document.paragraphs if p.text.strip()]

    assert any("Teknoloji" in t for t in paragraph_texts)
    template_default_headings = [
        "AMERIKA", "AVRUPA", "RUSYA / UKRAYNA", "TÜRKİYE",
        "KAFKASYA / ORTADOĞU", "AFRIKA", "ASYA", "TEKNOLOJI / YAPAY ZEKA",
    ]
    leaked = [h for h in template_default_headings if any(h in t for t in paragraph_texts)]
    assert leaked == []


def test_generate_bulletin_respects_display_order(client, auth_headers, db_session, monkeypatch):
    """The bundled template's Heading1 paragraphs sit in a fixed physical
    order (AMERIKA before AVRUPA). Regression test for a reported issue where
    the generated report's index/table-of-contents always followed that
    fixed template order instead of the user-configured display_order."""
    async def classify_by_display_order(db, articles, category_names, blocks=None):
        return [
            {"article_id": a.id, "top_category": category_names[i], "subcategory": "Alt Başlık", "type": "haber"}
            for i, a in enumerate(articles)
        ]

    monkeypatch.setattr(bulletin_service, "classify_articles_for_bulletin", classify_by_display_order)

    # AMERIKA precedes AVRUPA in the template's physical layout, but the
    # user has configured the opposite display_order here.
    _make_bulletin_category(db_session, name="AMERIKA", display_order=1)
    _make_bulletin_category(db_session, name="AVRUPA", display_order=0)
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Avrupa Haberi", published_at=now - timedelta(hours=1))
    _make_article(db_session, feed.id, title="Amerika Haberi", published_at=now - timedelta(hours=2))

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))
    heading1_texts = [p.text.strip() for p in document.paragraphs if p.style.style_id == "Heading1" and p.text.strip()]

    assert heading1_texts.index("AVRUPA") < heading1_texts.index("AMERIKA")


def test_generate_bulletin_case_variant_categories_get_separate_headings(
    client, auth_headers, db_session, monkeypatch
):
    """"Avrupa" and "AVRUPA" both fold to the same key under
    docx_service._normalize() (Turkish-I-aware case-fold). Categories this
    close can't be created via the admin API anymore (case-insensitive
    duplicate check), but this models a database that predates that guard by
    writing the row directly. Regression test for
    docx_service._rebuild_category_sections handing the SAME template
    heading paragraph to both categories (`existing_by_name.get()` never
    removed the entry) — the second category would relocate the first one's
    already-placed heading out from under its own content, stranding the
    first category's articles above an orphaned heading."""
    async def classify_by_category_names(db, articles, category_names, blocks=None):
        return [
            {"article_id": a.id, "top_category": category_names[i], "subcategory": "Alt Başlık", "type": "haber"}
            for i, a in enumerate(articles)
        ]

    monkeypatch.setattr(bulletin_service, "classify_articles_for_bulletin", classify_by_category_names)

    _make_bulletin_category(db_session, name="Avrupa", display_order=0)
    _make_bulletin_category(db_session, name="AVRUPA", display_order=1)
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Avrupa Haberi", published_at=now - timedelta(hours=1))
    _make_article(db_session, feed.id, title="Ikinci Avrupa Haberi", published_at=now - timedelta(hours=2))

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))

    heading1_avrupa_indices = [
        i for i, p in enumerate(document.paragraphs)
        if p.style.style_id == "Heading1" and p.text.strip().upper() == "AVRUPA"
    ]
    assert len(heading1_avrupa_indices) == 2, (
        "each case-variant category must get its own Heading1 paragraph, not "
        "share/relocate the other's"
    )

    paragraph_texts = [p.text for p in document.paragraphs]
    first_heading_idx, second_heading_idx = heading1_avrupa_indices
    first_article_idx = paragraph_texts.index("📌 Test Feed - Avrupa Haberi")
    second_article_idx = paragraph_texts.index("📌 Test Feed - Ikinci Avrupa Haberi")

    # Each category's article must sit under its OWN heading, not stranded
    # above it or under the other category's heading.
    assert first_heading_idx < first_article_idx < second_heading_idx < second_article_idx


def test_generate_bulletin_too_many_articles_returns_400(client, auth_headers, db_session, monkeypatch):
    monkeypatch.setattr(app_settings, "bulletin_max_articles", 2)

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    for i in range(3):
        _make_article(db_session, feed.id, title=f"Article {i}", published_at=now - timedelta(hours=1))

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 400


def test_bulletin_max_articles_defaults_to_50():
    from app.core.config import Settings

    assert Settings.model_fields["bulletin_max_articles"].default == 50


def test_generate_bulletin_uses_classification_cache(client, auth_headers, db_session, monkeypatch):
    call_count = {"n": 0}

    async def counting_classify(db, articles, category_names, blocks=None):
        if not articles:
            return []
        call_count["n"] += 1
        return [
            {"article_id": a.id, "top_category": category_names[0], "subcategory": "Test", "type": "haber"}
            for a in articles
        ]

    monkeypatch.setattr(bulletin_service, "classify_articles_for_bulletin", counting_classify)

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Cached Haber", published_at=now - timedelta(hours=1))

    first = client.post("/api/bulletin/generate", json={}, headers=auth_headers)
    assert first.status_code == 200
    assert call_count["n"] == 1

    second = client.post("/api/bulletin/generate", json={}, headers=auth_headers)
    assert second.status_code == 200
    assert call_count["n"] == 1  # second run served entirely from the classification cache


# ==================== POST /generate — favorites ====================

def test_generate_bulletin_include_favorites_only(client, auth_headers, db_session, monkeypatch):

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(
        db_session, feed.id, title="Favori Haber", is_starred=True,
        published_at=now - timedelta(hours=1),
    )
    _make_article(
        db_session, feed.id, title="Normal Haber", is_starred=False,
        published_at=now - timedelta(hours=1),
    )

    response = client.post(
        "/api/bulletin/generate", json={"include_favorites": True}, headers=auth_headers
    )

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))
    paragraph_texts = [p.text for p in document.paragraphs if p.text.strip()]
    assert any("Favori Haber" in t for t in paragraph_texts)
    assert not any("Normal Haber" in t for t in paragraph_texts)


def test_generate_bulletin_priorities_and_favorites_union(client, auth_headers, db_session, monkeypatch):

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(
        db_session, feed.id, title="Yuksek Oncelik", priority="high", is_starred=False,
        published_at=now - timedelta(hours=1),
    )
    _make_article(
        db_session, feed.id, title="Favori Dusuk Oncelik", priority="low", is_starred=True,
        published_at=now - timedelta(hours=1),
    )
    _make_article(
        db_session, feed.id, title="Orta Oncelik Disarida", priority="med", is_starred=False,
        published_at=now - timedelta(hours=1),
    )

    response = client.post(
        "/api/bulletin/generate",
        json={"priorities": ["high"], "include_favorites": True},
        headers=auth_headers,
    )

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))
    paragraph_texts = [p.text for p in document.paragraphs if p.text.strip()]
    assert any("Yuksek Oncelik" in t for t in paragraph_texts)
    assert any("Favori Dusuk Oncelik" in t for t in paragraph_texts)
    assert not any("Orta Oncelik Disarida" in t for t in paragraph_texts)


def test_generate_bulletin_dedup_when_article_matches_both(client, auth_headers, db_session, monkeypatch):
    """An article that is both a matching priority AND starred must appear
    exactly once in the bulletin, not twice."""

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(
        db_session, feed.id, title="Hem Yuksek Hem Favori", priority="high", is_starred=True,
        published_at=now - timedelta(hours=1),
    )

    response = client.post(
        "/api/bulletin/generate",
        json={"priorities": ["high"], "include_favorites": True},
        headers=auth_headers,
    )

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))
    paragraph_texts = [p.text for p in document.paragraphs if p.text.strip()]
    # A broken union (fetching the article via two separate queries instead of one OR'd query)
    # would show up as a duplicate article header.
    headers = [t for t in paragraph_texts if t.startswith("📌") and "Hem Yuksek Hem Favori" in t]
    assert len(headers) == 1


def test_generate_bulletin_include_favorites_defaults_false(client, auth_headers, db_session, monkeypatch):

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    # Starred but priority doesn't match anything requested, and the request
    # doesn't mention include_favorites at all (mirrors test_generate_bulletin_success).
    _make_article(
        db_session, feed.id, title="Favori Ama Istenmedi", priority="low", is_starred=True,
        published_at=now - timedelta(hours=1),
    )

    response = client.post(
        "/api/bulletin/generate", json={"priorities": ["high"]}, headers=auth_headers
    )

    assert response.status_code == 400  # nothing matches priority "high" and favorites weren't requested


# ==================== GET /preview-count ====================

def test_preview_count_requires_auth(client):
    response = client.get("/api/bulletin/preview-count")
    assert response.status_code == 401


def test_preview_count_defaults_to_last_24h_all_priorities(client, auth_headers, db_session):
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Icerde 1", published_at=now - timedelta(hours=1))
    _make_article(db_session, feed.id, title="Icerde 2", published_at=now - timedelta(hours=2))
    _make_article(db_session, feed.id, title="Disarda", published_at=now - timedelta(hours=48))

    response = client.get("/api/bulletin/preview-count", headers=auth_headers)

    assert response.status_code == 200
    assert response.json() == {"count": 2}


def test_preview_count_filters_by_priorities(client, auth_headers, db_session):
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, priority="high", published_at=now - timedelta(hours=1))
    _make_article(db_session, feed.id, priority="low", published_at=now - timedelta(hours=1))

    response = client.get(
        "/api/bulletin/preview-count", params={"priorities": "high"}, headers=auth_headers
    )

    assert response.status_code == 200
    assert response.json() == {"count": 1}


def test_preview_count_include_favorites_unions_with_priorities(client, auth_headers, db_session):
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, priority="high", is_starred=False, published_at=now - timedelta(hours=1))
    _make_article(db_session, feed.id, priority="low", is_starred=True, published_at=now - timedelta(hours=1))
    _make_article(db_session, feed.id, priority="med", is_starred=False, published_at=now - timedelta(hours=1))

    response = client.get(
        "/api/bulletin/preview-count",
        params={"priorities": "high", "include_favorites": "true"},
        headers=auth_headers,
    )

    assert response.status_code == 200
    assert response.json() == {"count": 2}


def test_preview_count_dedup_when_overlapping(client, auth_headers, db_session):
    """A naive 'count matches + count starred' implementation would double-count
    an article that satisfies both — the real union must count it once."""
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, priority="high", is_starred=True, published_at=now - timedelta(hours=1))

    response = client.get(
        "/api/bulletin/preview-count",
        params={"priorities": "high", "include_favorites": "true"},
        headers=auth_headers,
    )

    assert response.status_code == 200
    assert response.json() == {"count": 1}


def test_preview_count_rejects_invalid_priority(client, auth_headers):
    response = client.get(
        "/api/bulletin/preview-count", params={"priorities": "urgent"}, headers=auth_headers
    )
    assert response.status_code == 400


def test_preview_count_zero_when_no_matches(client, auth_headers, db_session):
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, priority="low", published_at=now - timedelta(hours=1))

    response = client.get(
        "/api/bulletin/preview-count", params={"priorities": "high"}, headers=auth_headers
    )

    assert response.status_code == 200
    assert response.json() == {"count": 0}


def test_preview_count_works_without_categories(client, auth_headers, db_session):
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, published_at=now - timedelta(hours=1))

    response = client.get("/api/bulletin/preview-count", headers=auth_headers)

    assert response.status_code == 200
    assert response.json() == {"count": 1}


# ==================== POST /generate — template artifact regressions ====================

def test_generate_bulletin_no_blank_heading1_paragraphs(client, auth_headers, db_session, monkeypatch):
    """The bundled template contains several empty Heading1 paragraphs used as
    layout spacers (a Google Docs export artifact). An earlier version of the
    stray-heading cleanup keyed template headings by normalized text in a
    dict, so multiple blank headings collided under the same '' key and all
    but one silently survived un-removed in the output — orphaned blank
    Heading1 paragraphs that broke Word's İçindekiler (Table of Contents)
    generation on open."""

    # A category name unrelated to the template's own defaults forces every
    # bundled Heading1 (including its blank spacers) through the stray-removal path.
    _make_bulletin_category(db_session, name="Teknoloji")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Yapay Zeka Haberi", published_at=now - timedelta(hours=1))

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 200
    document = DocxDocument(io.BytesIO(response.content))
    blank_heading1s = [
        p for p in document.paragraphs if p.style.style_id == "Heading1" and not p.text.strip()
    ]
    assert blank_heading1s == []


def test_generate_bulletin_toc_field_is_locale_independent(client, auth_headers, db_session, monkeypatch):
    """The bundled template's TOC field originally selected entries purely
    via \\t, a literal English-only style-name list ("Heading 1,1,Heading
    2,2,..."). \\t matches a paragraph by its LOCALIZED style display name,
    and Word does not translate "Heading 1" for that comparison on a
    non-English install (confirmed via Word COM automation on a German
    installation: the paragraph's outline level was correctly 1, but its
    style name is "Überschrift 1", so \\t alone matched nothing and the TOC
    came back with zero entries) — reproducing exactly the "no heading style
    applied" dialog reported against a real generated bulletin. \\o "1-6"
    additionally matches by outline level, which is locale-independent."""

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Test Haberi", published_at=now - timedelta(hours=1))

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8")
    toc_instr_start = document_xml.index("<w:instrText")
    toc_instr_end = document_xml.index("</w:instrText>", toc_instr_start)
    # instrText may be split across sibling elements; the TOC's own opening
    # instrText run is the first one in the document (inside the bundled
    # İçindekiler content control), so this is enough to check its content.
    toc_instr_text = document_xml[toc_instr_start:toc_instr_end]
    assert "TOC" in toc_instr_text
    assert '\\o "1-6"' in toc_instr_text or "\\o &quot;1-6&quot;" in toc_instr_text


# ==================== Generated Bulletin Persistence ====================

def test_generate_bulletin_persists_a_row(client, auth_headers, db_session, monkeypatch):

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Test Haberi", priority="high", published_at=now - timedelta(hours=1))

    response = client.post(
        "/api/bulletin/generate",
        json={"priorities": ["high"], "include_favorites": True},
        headers=auth_headers,
    )

    assert response.status_code == 200
    rows = crud.get_generated_bulletins(db_session)
    assert len(rows) == 1
    row = rows[0]
    assert row.article_count == 1
    assert row.priorities == "high"
    assert row.include_favorites is True
    assert Path(row.stored_path).exists()


def test_list_generated_bulletins_requires_auth(client):
    response = client.get("/api/bulletin/generated")
    assert response.status_code == 401


def test_list_generated_bulletins_returns_newest_first(client, auth_headers, db_session, monkeypatch):

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Birinci", published_at=now - timedelta(hours=1))
    client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    _make_article(db_session, feed.id, title="Ikinci", published_at=now - timedelta(hours=1))
    client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    response = client.get("/api/bulletin/generated", headers=auth_headers)

    assert response.status_code == 200
    body = response.json()
    assert len(body) == 2
    assert body[0]["article_count"] == 2  # most recent generate saw both articles
    assert body[1]["article_count"] == 1


def test_download_generated_bulletin_requires_auth(client):
    response = client.get("/api/bulletin/generated/1/download")
    assert response.status_code == 401


def test_download_generated_bulletin_404_when_missing(client, auth_headers):
    response = client.get("/api/bulletin/generated/999999/download", headers=auth_headers)
    assert response.status_code == 404


def test_download_generated_bulletin_returns_the_saved_file(client, auth_headers, db_session, monkeypatch):

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Test Haberi", published_at=now - timedelta(hours=1))
    client.post("/api/bulletin/generate", json={}, headers=auth_headers)
    bulletin_id = crud.get_generated_bulletins(db_session)[0].id

    response = client.get(f"/api/bulletin/generated/{bulletin_id}/download", headers=auth_headers)

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.content[:4] == b"PK\x03\x04"


def test_delete_generated_bulletin_requires_auth(client):
    response = client.delete("/api/bulletin/generated/1")
    assert response.status_code == 401


def test_delete_generated_bulletin_404_when_missing(client, admin_headers):
    response = client.delete("/api/bulletin/generated/999999", headers=admin_headers)
    assert response.status_code == 404


def test_delete_generated_bulletin_removes_row_and_file(client, admin_headers, db_session, monkeypatch):

    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, title="Test Haberi", published_at=now - timedelta(hours=1))
    client.post("/api/bulletin/generate", json={}, headers=admin_headers)
    row = crud.get_generated_bulletins(db_session)[0]
    stored_path = Path(row.stored_path)
    assert stored_path.exists()

    response = client.delete(f"/api/bulletin/generated/{row.id}", headers=admin_headers)

    assert response.status_code == 200
    assert crud.get_generated_bulletin(db_session, row.id) is None
    assert not stored_path.exists()


def test_generate_bulletin_survives_duplicate_and_string_article_ids(client, auth_headers, db_session, monkeypatch):
    """Regression (#29): the same article_id twice in the LLM answer hit the
    (article_id, category_set_hash) unique constraint and /generate returned 500; an id sent as a
    string ("123") silently dropped the article from the report."""
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    now = datetime.now(timezone.utc)
    article = _make_article(db_session, feed.id, title="Tekrarlanan Haber", priority="high",
                            published_at=now - timedelta(hours=1))

    async def duplicate_classify(db, articles, category_names, blocks=None):
        return [
            {"article_id": str(article.id), "top_category": "AVRUPA", "subcategory": "İlk", "type": "haber"},
            {"article_id": article.id, "top_category": "AVRUPA", "subcategory": "İkinci", "type": "haber"},
            {"article_id": article.id, "top_category": "AVRUPA", "subcategory": "Üçüncü", "type": "haber"},
        ]

    monkeypatch.setattr(bulletin_service, "classify_articles_for_bulletin", duplicate_classify)

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    assert response.status_code == 200
    rows = db_session.query(models.ArticleBulletinClassification).all()
    assert [(r.article_id, r.subcategory) for r in rows] == [(article.id, "İlk")]
    texts = [p.text for p in DocxDocument(io.BytesIO(response.content)).paragraphs]
    assert texts.count("📌 Test Feed - Tekrarlanan Haber") == 1


def test_generate_bulletin_names_the_feed_as_source(client, auth_headers, db_session, monkeypatch):
    """Regression (#31): cards and summaries show the feed name as the source, the Word bulletin
    still showed the article URL's hostname."""
    _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session, title="Sputnik Türkiye")
    now = datetime.now(timezone.utc)
    _make_article(db_session, feed.id, url="https://tr.sputniknews.example/haber-1", priority="high",
                  published_at=now - timedelta(hours=1))

    response = client.post("/api/bulletin/generate", json={}, headers=auth_headers)

    body_text = "\n".join(p.text for p in DocxDocument(io.BytesIO(response.content)).paragraphs)
    assert "Sputnik Türkiye" in body_text
    assert "tr.sputniknews.example" not in body_text


def test_classification_matches_category_names_case_and_turkish_i_insensitively(monkeypatch, db_session):
    """Regression: user categories are stored as typed ("Ukrayna", "Amerika", "Iran") but the model
    answers in upper case ("UKRAYNA", "AMERİKA", "İRAN"). The exact-match check sent every such
    article to DİĞER, so the bulletin had Ukraine/America news under DİĞER next to empty-ish
    Ukrayna/Amerika sections."""
    async def fake_completion(system_prompt, user_prompt, max_completion_tokens):
        return {"classifications": [
            {"article_id": 1, "top_category": "UKRAYNA", "type": "haber"},
            {"article_id": 2, "top_category": "AMERİKA", "type": "haber"},
            {"article_id": 3, "top_category": "İRAN", "type": "haber"},
            {"article_id": 4, "top_category": " nato ", "type": "haber"},
            {"article_id": 5, "top_category": "ASYA", "type": "haber"},
        ]}

    async def no_limits():
        return None

    monkeypatch.setattr(bulletin_service, "_call_json_completion", fake_completion)
    monkeypatch.setattr(bulletin_service, "check_cost_limits", no_limits)
    articles = [models.Article(id=i, title=f"Haber {i}", url=f"https://x/{i}") for i in range(1, 6)]

    result = asyncio.run(_real_classify_articles_for_bulletin(
        db_session, articles, ["Türkiye", "Amerika", "Avrupa", "Ukrayna", "Iran", "Nato"]
    ))

    assert [r["top_category"] for r in result] == ["Ukrayna", "Amerika", "Iran", "Nato", "DİĞER"]


def test_generate_bulletin_places_cached_case_variant_category_under_its_section(
    client, auth_headers, db_session, monkeypatch
):
    """A cached classification whose top_category differs from the current category name only by
    case (e.g. the category was renamed "Avrupa" -> "AVRUPA"; the cache key ignores case) must land
    in that category — it used to fall into a bucket that was never rendered, dropping the article."""
    category = _make_bulletin_category(db_session, name="AVRUPA")
    feed = _make_feed(db_session)
    article = _make_article(db_session, feed.id, title="Brüksel Haberi",
                            published_at=datetime.now(timezone.utc) - timedelta(hours=1))
    crud.create_bulletin_classification(
        db_session, article_id=article.id,
        category_set_hash=bulletin_service.compute_category_set_hash([category.name]),
        top_category="Avrupa", subcategory="", article_type="haber", model_used="gpt-test",
    )

    texts = [p.text for p in _paragraphs(_generate(client, auth_headers))]

    assert "📌 Test Feed - Brüksel Haberi" in texts
    assert texts.index("AVRUPA") < texts.index("📌 Test Feed - Brüksel Haberi")
    assert "DİĞER" not in texts
